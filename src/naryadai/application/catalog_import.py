"""Bounded, add-only XLSX import validation for equipment and materials."""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass
from typing import Literal, NoReturn

import xlsxwriter  # type: ignore[import-untyped]
from anyio import CapacityLimiter, to_thread
from defusedxml import ElementTree
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from naryadai.application.common import OperationError, ensure_role
from naryadai.auth.dependencies import Principal
from naryadai.infrastructure.database import Database
from naryadai.infrastructure.models import Area, Equipment, Material

MAX_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_BYTES = 10 * 1024 * 1024
MAX_ZIP_ENTRIES = 100
MAX_ROWS = 500
MAX_COLUMNS = 20
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_EQUIPMENT_SHEETS = {"Оборудование", "Equipment"}
_MATERIAL_SHEETS = {"Материалы", "Materials"}
_REFERENCE_SHEETS = {"Справочник участков", "Areas"}
_EQUIPMENT_HEADERS = {
    "area_code": "area_code",
    "код участка": "area_code",
    "inventory_number": "inventory_number",
    "инвентарный номер": "inventory_number",
    "name": "name",
    "наименование": "name",
    "equipment_type": "equipment_type",
    "тип оборудования": "equipment_type",
    "criticality": "criticality",
    "критичность": "criticality",
}
_MATERIAL_HEADERS = {
    "code": "code",
    "код": "code",
    "name": "name",
    "наименование": "name",
    "unit": "unit",
    "единица измерения": "unit",
}


def _column_number(letters: str) -> int:
    number = 0
    for letter in letters:
        number = number * 26 + ord(letter.upper()) - 64
    return number


class ImportIssue(BaseModel):
    model_config = ConfigDict(frozen=True)

    field: str | None = None
    code: str
    message: str


class ImportRow(BaseModel):
    model_config = ConfigDict(frozen=True)

    sheet: Literal["equipment", "materials"]
    row: int = Field(ge=2)
    key: str
    values: dict[str, str | int]
    status: Literal["create", "skip", "error"]
    errors: list[ImportIssue] = Field(default_factory=list)


class ImportSection(BaseModel):
    model_config = ConfigDict(frozen=True)

    total: int = 0
    create: int = 0
    skip: int = 0
    error: int = 0
    rows: list[ImportRow] = Field(default_factory=list)


class CatalogImportPreview(BaseModel):
    model_config = ConfigDict(frozen=True)

    file_sha256: str
    fingerprint: str
    equipment: ImportSection
    materials: ImportSection
    created_equipment: int = 0
    created_materials: int = 0
    notes: list[str] = Field(default_factory=list)


@dataclass(frozen=True)
class _EquipmentInput:
    row: int
    area_code: str
    inventory_number: str
    name: str
    equipment_type: str
    criticality: int


@dataclass(frozen=True)
class _MaterialInput:
    row: int
    code: str
    name: str
    unit: str


@dataclass(frozen=True)
class ParsedCatalog:
    equipment: tuple[_EquipmentInput, ...]
    materials: tuple[_MaterialInput, ...]
    equipment_rows: tuple[ImportRow, ...]
    material_rows: tuple[ImportRow, ...]
    digest: str


def _fatal(code: str) -> NoReturn:
    raise OperationError(422, code)


def _text(value: object, *, maximum: int) -> str | None:
    if not isinstance(value, str):
        return None
    normalized = " ".join(value.strip().split())
    return normalized if normalized and len(normalized) <= maximum else None


def _headers(
    row: tuple[object, ...], mapping: dict[str, str], expected: set[str]
) -> list[str] | None:
    values: list[str] = []
    for value in row:
        text = _text(value, maximum=80)
        if text is None:
            return None
        key = mapping.get(text.casefold())
        if key is None or key in values:
            return None
        values.append(key)
    return values if set(values) == expected else None


def _safe_xlsx(content: bytes) -> None:
    if not content or len(content) > MAX_BYTES or not zipfile.is_zipfile(io.BytesIO(content)):
        _fatal("invalid_import_file")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if (
            len(entries) > MAX_ZIP_ENTRIES
            or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES
        ):
            _fatal("import_archive_too_large")
        names: set[str] = set()
        for item in entries:
            name = item.filename.lower()
            if (
                name in names
                or name.startswith(("/", "\\"))
                or ".." in name.split("/")
                or "\\" in name
                or item.flag_bits & 0x1
                or name.endswith("vbaproject.bin")
                or "externallinks/" in name
                or "/activex/" in name
                or "/ctrlprops/" in name
                or "/connections" in name
            ):
                _fatal("unsafe_import_workbook")
            names.add(name)
            payload = archive.read(item)
            if name.endswith((".xml", ".rels")):
                try:
                    root = ElementTree.fromstring(payload)
                except Exception as error:
                    raise OperationError(422, "invalid_import_workbook") from error
                for node in root.iter():
                    local = node.tag.rsplit("}", 1)[-1]
                    attributes = {
                        key.rsplit("}", 1)[-1].casefold(): value
                        for key, value in node.attrib.items()
                    }
                    if (
                        name.endswith(".rels")
                        and attributes.get("targetmode", "").casefold() == "external"
                    ):
                        _fatal("unsafe_import_workbook")
                    if "/worksheets/" in name and local == "f":
                        _fatal("unsafe_import_workbook")
                    if "/worksheets/" in name and local == "row":
                        row_reference = attributes.get("r", "")
                        if not row_reference.isdecimal() or int(row_reference) > MAX_ROWS + 1:
                            _fatal("import_sheet_limit_exceeded")
                    if "/worksheets/" in name and local == "c":
                        reference = attributes.get("r", "")
                        match = re.fullmatch(r"([A-Za-z]+)([0-9]+)", reference)
                        if (
                            match is None
                            or _column_number(match.group(1)) > MAX_COLUMNS
                            or int(match.group(2)) > MAX_ROWS + 1
                        ):
                            _fatal("import_sheet_limit_exceeded")


def parse_workbook(content: bytes) -> ParsedCatalog:
    """CPU-bound parser; call through a bounded worker thread."""

    try:
        _safe_xlsx(content)
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except OperationError:
        raise
    except Exception as error:
        raise OperationError(422, "invalid_import_workbook") from error
    try:
        return _parse_open_book(book, content)
    except OperationError:
        raise
    except Exception as error:
        raise OperationError(422, "invalid_import_workbook") from error
    finally:
        book.close()


def _parse_open_book(book: Workbook, content: bytes) -> ParsedCatalog:
    equipment_inputs: list[_EquipmentInput] = []
    material_inputs: list[_MaterialInput] = []
    equipment_rows: list[ImportRow] = []
    material_rows: list[ImportRow] = []
    names = set(book.sheetnames)
    equipment_name = next((name for name in _EQUIPMENT_SHEETS if name in names), None)
    material_name = next((name for name in _MATERIAL_SHEETS if name in names), None)
    if (
        equipment_name is None
        or material_name is None
        or len(names & _EQUIPMENT_SHEETS) != 1
        or len(names & _MATERIAL_SHEETS) != 1
        or len(names - _REFERENCE_SHEETS - _EQUIPMENT_SHEETS - _MATERIAL_SHEETS)
    ):
        _fatal("invalid_import_sheets")

    def parse_rows(sheet_name: str, kind: Literal["equipment", "materials"]) -> None:
        sheet = book[sheet_name]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=False)
        header_cells = next(rows, None)
        if header_cells is None:
            _fatal("invalid_import_headers")
        raw_header = tuple(cell.value for cell in header_cells)
        header = _headers(
            raw_header,
            _EQUIPMENT_HEADERS if kind == "equipment" else _MATERIAL_HEADERS,
            {"area_code", "inventory_number", "name", "equipment_type", "criticality"}
            if kind == "equipment"
            else {"code", "name", "unit"},
        )
        if header is None:
            _fatal("invalid_import_headers")
        for row_number, cells in enumerate(rows, start=2):
            if row_number > MAX_ROWS + 1 or len(cells) > MAX_COLUMNS:
                _fatal("import_sheet_limit_exceeded")
            if any(cell.data_type == "f" for cell in cells):
                _fatal("import_formulas_not_allowed")
            if any(cell.value not in (None, "") for cell in cells[len(header) :]):
                _fatal("invalid_import_headers")
            values = {
                field: cells[index].value if index < len(cells) else None
                for index, field in enumerate(header)
            }
            if not any(value not in (None, "") for value in values.values()):
                continue
            issues: list[ImportIssue] = []
            if kind == "equipment":
                area_code = _text(values["area_code"], maximum=32)
                inventory = _text(values["inventory_number"], maximum=64)
                name = _text(values["name"], maximum=160)
                equipment_type = _text(values["equipment_type"], maximum=64)
                criticality = values["criticality"]
                if area_code is None:
                    issues.append(
                        ImportIssue(
                            field="area_code", code="invalid", message="Укажите код участка."
                        )
                    )
                if inventory is None:
                    issues.append(
                        ImportIssue(
                            field="inventory_number",
                            code="invalid",
                            message="Укажите инвентарный номер.",
                        )
                    )
                if name is None:
                    issues.append(
                        ImportIssue(field="name", code="invalid", message="Укажите наименование.")
                    )
                if equipment_type is None:
                    issues.append(
                        ImportIssue(
                            field="equipment_type",
                            code="invalid",
                            message="Укажите тип оборудования.",
                        )
                    )
                if (
                    isinstance(criticality, bool)
                    or not isinstance(criticality, int)
                    or not 1 <= criticality <= 5
                ):
                    issues.append(
                        ImportIssue(
                            field="criticality",
                            code="invalid",
                            message="Критичность: целое число от 1 до 5.",
                        )
                    )
                key = inventory or f"row-{row_number}"
                display_values = {
                    field: value
                    if field == "criticality"
                    and isinstance(value, int)
                    and not isinstance(value, bool)
                    else str(value)
                    for field, value in values.items()
                    if value is not None
                }
                row = ImportRow(
                    sheet=kind,
                    row=row_number,
                    key=key,
                    values=display_values,
                    status="error" if issues else "create",
                    errors=issues,
                )
                equipment_rows.append(row)
                if not issues:
                    assert area_code is not None and inventory is not None
                    assert name is not None and equipment_type is not None
                    assert isinstance(criticality, int)
                    equipment_inputs.append(
                        _EquipmentInput(
                            row_number, area_code, inventory, name, equipment_type, criticality
                        )
                    )
            else:
                code = _text(values["code"], maximum=32)
                name = _text(values["name"], maximum=160)
                unit = _text(values["unit"], maximum=24)
                if code is not None and not re.fullmatch(r"[A-Za-z0-9_-]{2,32}", code):
                    code = None
                if code is None:
                    issues.append(
                        ImportIssue(field="code", code="invalid", message="Укажите код материала.")
                    )
                if name is None:
                    issues.append(
                        ImportIssue(field="name", code="invalid", message="Укажите наименование.")
                    )
                if unit is None:
                    issues.append(
                        ImportIssue(
                            field="unit", code="invalid", message="Укажите единицу измерения."
                        )
                    )
                key = code or f"row-{row_number}"
                row = ImportRow(
                    sheet=kind,
                    row=row_number,
                    key=key,
                    values={k: str(v) for k, v in values.items() if v is not None},
                    status="error" if issues else "create",
                    errors=issues,
                )
                material_rows.append(row)
                if not issues:
                    assert code is not None and name is not None and unit is not None
                    material_inputs.append(_MaterialInput(row_number, code, name, unit))

    parse_rows(equipment_name, "equipment")
    parse_rows(material_name, "materials")
    if len(equipment_rows) + len(material_rows) > MAX_ROWS:
        _fatal("import_row_limit_exceeded")
    return ParsedCatalog(
        tuple(equipment_inputs),
        tuple(material_inputs),
        tuple(equipment_rows),
        tuple(material_rows),
        hashlib.sha256(content).hexdigest(),
    )


def _section(rows: list[ImportRow]) -> ImportSection:
    return ImportSection(
        total=len(rows),
        create=sum(row.status == "create" for row in rows),
        skip=sum(row.status == "skip" for row in rows),
        error=sum(row.status == "error" for row in rows),
        rows=rows,
    )


async def preview_import(
    database: Database, principal: Principal, parsed: ParsedCatalog
) -> CatalogImportPreview:
    ensure_role(principal, "admin")
    async with database.sessions() as session:
        return await _preview_in_session(session, parsed)


async def _preview_in_session(session: AsyncSession, parsed: ParsedCatalog) -> CatalogImportPreview:
    areas = {row.code: row for row in await session.scalars(select(Area))}
    equipment = {row.inventory_number: row for row in await session.scalars(select(Equipment))}
    materials = {row.code: row for row in await session.scalars(select(Material))}
    equipment_rows = list(parsed.equipment_rows)
    material_rows = list(parsed.material_rows)
    equipment_index = {row.row: index for index, row in enumerate(equipment_rows)}
    material_index = {row.row: index for index, row in enumerate(material_rows)}
    seen_equipment: set[str] = set()
    seen_materials: set[str] = set()
    for item in parsed.equipment:
        index = equipment_index[item.row]
        row = equipment_rows[index]
        issues = list(row.errors)
        key = item.inventory_number
        area = areas.get(item.area_code)
        if key in seen_equipment:
            issues.append(
                ImportIssue(
                    field="inventory_number",
                    code="duplicate",
                    message="Дубликат инвентарного номера в файле.",
                )
            )
        seen_equipment.add(key)
        if area is None or not area.is_active:
            issues.append(
                ImportIssue(
                    field="area_code", code="unknown_area", message="Активный участок не найден."
                )
            )
        existing = equipment.get(key)
        status: Literal["create", "skip", "error"] = "create"
        if existing is not None:
            status = (
                "skip"
                if (existing.area_id == area.id if area else False)
                and existing.name == item.name
                and existing.equipment_type == item.equipment_type
                and existing.criticality == item.criticality
                else "error"
            )
            if status == "error":
                issues.append(
                    ImportIssue(
                        field="inventory_number",
                        code="conflict",
                        message="Инвентарный номер уже занят другими данными.",
                    )
                )
        equipment_rows[index] = row.model_copy(
            update={"status": "error" if issues else status, "errors": issues}
        )
    for material_item in parsed.materials:
        index = material_index[material_item.row]
        row = material_rows[index]
        issues = list(row.errors)
        key = material_item.code
        if key in seen_materials:
            issues.append(
                ImportIssue(field="code", code="duplicate", message="Дубликат кода в файле.")
            )
        seen_materials.add(key)
        existing_material = materials.get(key)
        material_status: Literal["create", "skip", "error"] = "create"
        if existing_material is not None:
            material_status = (
                "skip"
                if existing_material.name == material_item.name
                and existing_material.unit == material_item.unit
                else "error"
            )
            if material_status == "error":
                issues.append(
                    ImportIssue(
                        field="code",
                        code="conflict",
                        message="Код материала уже занят другими данными.",
                    )
                )
        material_rows[index] = row.model_copy(
            update={"status": "error" if issues else material_status, "errors": issues}
        )
    snapshot = {
        "areas": sorted((row.code, str(row.id), row.is_active) for row in areas.values()),
        "equipment": sorted(
            (row.inventory_number, str(row.area_id), row.name, row.equipment_type, row.criticality)
            for row in equipment.values()
        ),
        "materials": sorted((row.code, row.name, row.unit) for row in materials.values()),
    }
    canonical = {
        "digest": parsed.digest,
        "equipment": [row.model_dump() for row in equipment_rows],
        "materials": [row.model_dump() for row in material_rows],
        "snapshot": snapshot,
    }
    fingerprint = hashlib.sha256(
        json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    return CatalogImportPreview(
        file_sha256=parsed.digest,
        fingerprint=fingerprint,
        equipment=_section(equipment_rows),
        materials=_section(material_rows),
    )


async def apply_import(
    database: Database, principal: Principal, parsed: ParsedCatalog, fingerprint: str
) -> CatalogImportPreview:
    ensure_role(principal, "admin")
    async with database.sessions.begin() as session:
        # Also serialize with existing CRUD's SELECT FOR UPDATE before its
        # eventual write. Plain catalog reads remain available during import.
        await session.execute(text("LOCK TABLE areas, equipment, materials IN EXCLUSIVE MODE"))
        preview = await _preview_in_session(session, parsed)
        if preview.fingerprint != fingerprint:
            if not (
                preview.equipment.error == 0
                and preview.materials.error == 0
                and preview.equipment.create == 0
                and preview.materials.create == 0
            ):
                raise OperationError(409, "catalog_import_stale")
            return preview
        if preview.equipment.error or preview.materials.error:
            raise OperationError(422, "catalog_import_invalid")
        for item, row in zip(parsed.equipment, preview.equipment.rows, strict=True):
            if row.status == "create":
                area = await session.scalar(
                    select(Area).where(Area.code == item.area_code).with_for_update()
                )
                if area is None or not area.is_active:
                    raise OperationError(409, "catalog_import_stale")
                session.add(
                    Equipment(
                        area_id=area.id,
                        inventory_number=item.inventory_number,
                        name=item.name,
                        equipment_type=item.equipment_type,
                        criticality=item.criticality,
                    )
                )
        for material_item, row in zip(parsed.materials, preview.materials.rows, strict=True):
            if row.status == "create":
                session.add(
                    Material(
                        code=material_item.code, name=material_item.name, unit=material_item.unit
                    )
                )
        try:
            await session.flush()
        except IntegrityError as error:
            raise OperationError(409, "catalog_import_conflict") from error
    return preview.model_copy(
        update={
            "created_equipment": preview.equipment.create,
            "created_materials": preview.materials.create,
        }
    )


async def template_bytes(
    database: Database, principal: Principal, limiter: CapacityLimiter
) -> bytes:
    ensure_role(principal, "admin")
    async with database.sessions() as session:
        areas = list(await session.scalars(select(Area).where(Area.is_active).order_by(Area.code)))
    return await to_thread.run_sync(
        _build_template, [(area.code, area.name) for area in areas], limiter=limiter
    )


def _build_template(areas: list[tuple[str, str]]) -> bytes:
    stream = io.BytesIO()
    book = xlsxwriter.Workbook(
        stream,
        {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False},
    )
    header = book.add_format({"bold": True, "bg_color": "#D9EAF7"})
    identifier_format = book.add_format({"num_format": "@"})
    equipment = book.add_worksheet("Оборудование")
    materials = book.add_worksheet("Материалы")
    reference = book.add_worksheet("Справочник участков")
    equipment.write_row(
        0,
        0,
        ["код участка", "инвентарный номер", "наименование", "тип оборудования", "критичность"],
        header,
    )
    materials.write_row(0, 0, ["код", "наименование", "единица измерения"], header)
    reference.write_row(0, 0, ["код участка", "наименование"], header)
    for index, area in enumerate(areas, start=1):
        reference.write_row(index, 0, area)
    for sheet in (equipment, materials, reference):
        sheet.freeze_panes(1, 0)
        sheet.set_column(0, 4, 24)
    equipment.set_column(0, 3, 24, identifier_format)
    materials.set_column(0, 2, 24, identifier_format)
    equipment.data_validation(
        1,
        4,
        MAX_ROWS,
        4,
        {
            "validate": "integer",
            "criteria": "between",
            "minimum": 1,
            "maximum": 5,
            "input_title": "Критичность",
            "input_message": "Целое число от 1 до 5.",
            "error_title": "Проверьте значение",
            "error_message": "Укажите число от 1 до 5.",
        },
    )
    book.close()
    return stream.getvalue()
