"""Regression evidence for shared load, private push and unduplicated reports."""

import asyncio
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from anyio import CapacityLimiter
from openpyxl import load_workbook
from sqlalchemy import select
from test_brigade_assignments import _body, _team
from test_push_delivery import _settings

from naryadai.analytics.contracts import AnalyticsQuery
from naryadai.analytics.service import analytics_options, build_report
from naryadai.application import notification_engine, push_delivery
from naryadai.application.common import OperationError
from naryadai.application.contracts import Completion, MaterialLine, OrderAction
from naryadai.application.orders import apply_review, create_order, execute_action
from naryadai.application.queries import brigade_options, workload
from naryadai.application.report_evidence import DowntimeRecord, save_report_evidence
from naryadai.application.suggestions import SuggestionRequest, work_order_suggestions
from naryadai.domain.lifecycle import AiAssessment
from naryadai.infrastructure.models import (
    AuthSession,
    Brigade,
    Employee,
    Material,
    Notification,
    PushDelivery,
    PushSubscription,
    WorkOrder,
)
from naryadai.infrastructure.photo_store import PhotoStore
from naryadai.infrastructure.webpush import PushResult
from naryadai.reporting.order_report import export_order

pytestmark = pytest.mark.asyncio


async def test_overlapping_crews_cannot_start_concurrently_and_never_appear_free(database):
    data = await _team(database)
    jobs = []
    for index, responsible in enumerate((data["first"], data["second"])):
        issued = await create_order(
            database,
            data["master"],
            _body(data).model_copy(update={"executor_id": responsible.employee_id}),
            f"parallel-issue-{index}",
        )
        order_id = UUID(issued["order_id"])
        accepted = await execute_action(
            database,
            responsible,
            order_id,
            OrderAction(action="accept", expected_version=1),
            f"parallel-accept-{index}",
        )
        jobs.append((responsible, order_id, accepted["version"]))
    outcomes = await asyncio.wait_for(
        asyncio.gather(
            *(
                execute_action(
                    database,
                    actor,
                    order_id,
                    OrderAction(action="start", expected_version=version),
                    f"parallel-start-{index}",
                )
                for index, (actor, order_id, version) in enumerate(jobs)
            ),
            return_exceptions=True,
        ),
        timeout=10,
    )
    assert sum(isinstance(outcome, dict) for outcome in outcomes) == 1
    errors = [outcome for outcome in outcomes if isinstance(outcome, OperationError)]
    assert len(errors) == 1 and "participant_busy" in str(errors[0])
    people = await workload(database, data["master"])
    assert len(people) == 2
    assert all(person.availability == "busy" for person in people)
    assert len({person.current_order_id for person in people}) == 1
    options = await brigade_options(database, data["master"], data["area"])
    assert all(person.availability == "busy" for person in options.items[0].participants)
    suggestions = await work_order_suggestions(
        database,
        data["master"],
        SuggestionRequest(
            area_id=data["area"], equipment_id=data["equipment"], description="Repair"
        ),
    )
    assert len(suggestions.executors) == 2
    assert all(person.availability == "busy" for person in suggestions.executors)
    async with database.sessions() as session:
        order = await session.get(WorkOrder, jobs[0][1])
        assert await notification_engine._candidate(session, order, "m", 4) is None


async def test_participant_push_fanout_deduplicates_and_rechecks_assignment(database, monkeypatch):
    data = await _team(database)
    now = datetime.now(UTC)
    async with database.sessions.begin() as session:
        auth = AuthSession(
            employee_id=data["second"].employee_id,
            token_hash="p" * 64,
            created_at=now,
            expires_at=now + timedelta(hours=1),
        )
        session.add(auth)
        await session.flush()
        subscription = PushSubscription(
            employee_id=data["second"].employee_id,
            session_id=auth.id,
            endpoint="https://fcm.googleapis.com/fcm/send/crew",
            endpoint_hash="crew",
            p256dh="p",
            auth="a",
        )
        session.add(subscription)
        await session.flush()
        subscription_id = subscription.id
    orders = [
        await create_order(database, data["master"], _body(data), f"push-crew-issue-{index}")
        for index in range(2)
    ]
    moment = datetime.now(UTC)
    assert await notification_engine.process_outbox(database, now=moment) == 4
    assert await notification_engine.process_outbox(database, now=moment) == 0
    async with database.sessions() as session:
        notices = list(await session.scalars(select(Notification)))
        assert len(notices) == 4
        assert {notice.employee_id for notice in notices} == {
            data["first"].employee_id,
            data["second"].employee_id,
        }
    calls = []

    async def accepted(*args, **kwargs):
        calls.append((args, kwargs))
        return PushResult(201, None, None)

    monkeypatch.setattr(push_delivery, "send_web_push", accepted)
    assert await push_delivery.deliver_push(database, settings=_settings(), now=moment) == 2
    assert len(calls) == 2
    removed_order = UUID(orders[0]["order_id"])
    await execute_action(
        database,
        data["master"],
        removed_order,
        OrderAction(
            action="reassign",
            expected_version=1,
            executor_id=data["first"].employee_id,
            reason="One person now responsible",
        ),
        "push-remove-peer",
    )
    # Another current crew assignment and other employees' notifications must
    # not authorize a delayed private delivery for the removed assignment.
    async with database.sessions.begin() as session:
        notice = Notification(
            employee_id=data["second"].employee_id,
            work_order_id=removed_order,
            dedup_key="old-peer-delayed",
            kind="test",
            title="Private",
            body="Repair",
            payload={},
        )
        session.add(notice)
        await session.flush()
        delivery = PushDelivery(
            notification_id=notice.id,
            subscription_id=subscription_id,
            next_attempt_at=moment,
        )
        session.add(delivery)
        await session.flush()
        delivery_id = delivery.id
    assert await push_delivery.deliver_push(database, settings=_settings(), now=moment) == 0
    assert len(calls) == 2
    async with database.sessions() as session:
        delivery = await session.get(PushDelivery, delivery_id)
        assert delivery.failed_at is not None
        assert delivery.last_error == "push_recipient_ineligible"


async def _close_with_material(database, data, material_id, *, group):
    body = _body(data)
    if not group:
        body = body.model_copy(update={"brigade_id": None, "participant_ids": ()})
    result = await create_order(database, data["master"], body, f"report-issue-{group}")
    order_id = UUID(result["order_id"])
    for action in ("accept", "start", "complete"):
        completion = (
            Completion(
                work_description="Replaced seal and checked pressure.",
                fault_code_id=data["fault"],
                materials=(MaterialLine(material_id=material_id, quantity=Decimal(2)),),
            )
            if action == "complete"
            else None
        )
        result = await execute_action(
            database,
            data["first"],
            order_id,
            OrderAction(action=action, expected_version=result["version"], completion=completion),
            f"report-{action}-{group}",
        )
    result = await apply_review(
        database,
        order_id,
        expected_order_version=result["version"],
        submission_version=result["version"],
        verdict=AiAssessment.ACCEPTED,
        needs_master_review=False,
        score=5,
        explanation="Accepted",
        model_name="test",
    )
    result = await execute_action(
        database,
        data["master"],
        order_id,
        OrderAction(action="close", expected_version=result["version"]),
        f"report-close-{group}",
    )
    return order_id, result["version"]


async def test_team_metrics_count_once_preserve_legacy_and_export_human_names(database, tmp_path):
    data = await _team(database)
    async with database.sessions.begin() as session:
        material = Material(code="SEAL", name="Seal", unit="pc")
        session.add(material)
        await session.flush()
        material_id = material.id
    order_id, version = await _close_with_material(database, data, material_id, group=True)
    await _close_with_material(database, data, material_id, group=False)
    moment = datetime.now(UTC)
    await save_report_evidence(
        database,
        data["master"],
        order_id,
        DowntimeRecord(
            expected_version=version,
            started_at=moment - timedelta(hours=2),
            ended_at=moment - timedelta(hours=1),
            reason="Machine stopped for repair",
        ),
        "team-downtime-once",
    )
    query = AnalyticsQuery(
        period="custom",
        **{"from": moment - timedelta(days=1), "to": moment + timedelta(days=1)},
        timezone="UTC",
        brigade_ids=(data["brigade"],),
    )
    report = await build_report(database, data["master"], query)
    assert report.orders.issued == report.orders.closed == 2
    assert report.materials.usage[0].quantity == 4
    assert report.materials.usage[0].order_count == 2
    assert report.downtime.known_seconds == 3600
    assert report.ratings.brigades[0].sample_size == 2
    async with database.sessions.begin() as session:
        replacement = Brigade(code="BR-2", name="New brigade")
        session.add(replacement)
        await session.flush()
        replacement_id = replacement.id
        for person in (data["first"], data["second"]):
            employee = await session.get(Employee, person.employee_id)
            employee.brigade_id = replacement_id
    options = await analytics_options(database, data["master"])
    assert {item.id for item in options.brigades} == {data["brigade"], replacement_id}
    historical = await build_report(database, data["master"], query)
    legacy = await build_report(
        database,
        data["master"],
        query.model_copy(update={"brigade_ids": (replacement_id,)}),
    )
    assert historical.orders.closed == legacy.orders.closed == 1
    assert historical.materials.usage[0].quantity == legacy.materials.usage[0].quantity == 2
    assert historical.downtime.known_seconds == 3600
    assert legacy.downtime.known_seconds == 0
    store = PhotoStore(
        tmp_path,
        max_bytes=1_000_000,
        max_pixels=1_000_000,
        max_dimension=1600,
        output_max_bytes=1_000_000,
    )
    raw = await export_order(database, data["master"], order_id, store, CapacityLimiter(1))
    book = load_workbook(BytesIO(raw), read_only=True)
    try:
        cells = dict(book["Наряд"].iter_rows(min_row=2, values_only=True))
        assert cells["Бригада"] == "BR-1 — Brigade one"
        assert cells["Ответственный"] == "One"
        assert "One (ответственный)" in cells["Участники бригады"]
        assert "Two (участник)" in cells["Участники бригады"]
    finally:
        book.close()
