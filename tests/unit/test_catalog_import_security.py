"""Hostile workbook and boundary checks independent of the import implementation."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import xlsxwriter
from openpyxl import load_workbook

from naryadai.application.catalog_import import _build_template, parse_workbook
from naryadai.application.common import OperationError


def workbook(*, criticality=3, code="MAT-1", name="Seal", count=1):
    stream = BytesIO()
    book = xlsxwriter.Workbook(stream, {"in_memory": True})
    equipment = book.add_worksheet("Equipment")
    equipment.write_row(
        0, 0, ["area_code", "inventory_number", "name", "equipment_type", "criticality"]
    )
    equipment.write_row(1, 0, ["AREA", "00123", "Pump", "pump", criticality])
    materials = book.add_worksheet("Materials")
    materials.write_row(0, 0, ["code", "name", "unit"])
    for row in range(1, count + 1):
        materials.write_row(row, 0, [code if count == 1 else f"MAT-{row}", name, "pcs"])
    book.close()
    return stream.getvalue()


def replace_entry(content, path, transform):
    target = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(target, "w", ZIP_DEFLATED) as result:
        for item in source.infolist():
            payload = source.read(item)
            result.writestr(item, transform(payload) if item.filename == path else payload)
    return target.getvalue()


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
@pytest.mark.parametrize("reference", ["XFD1", "A1048576", "&#88;FD1"])
def test_sparse_xml_references_rejected_before_openpyxl(encoding, reference):
    def change(payload):
        text = payload.decode().replace('r="A1"', f"r='{reference}'", 1)
        return text.replace('encoding="UTF-8"', f'encoding="{encoding}"').encode(encoding)

    content = replace_entry(workbook(), "xl/worksheets/sheet1.xml", change)
    with pytest.raises(OperationError) as caught:
        parse_workbook(content)
    assert caught.value.status_code == 422


@pytest.mark.parametrize("criticality", [2.5, True, "high", -1, 6])
def test_invalid_criticality_is_row_error_not_server_exception(criticality):
    result = parse_workbook(workbook(criticality=criticality))
    assert result.equipment_rows[0].status == "error"
    assert any(issue.field == "criticality" for issue in result.equipment_rows[0].errors)


@pytest.mark.parametrize("code", ["X", "bad code", "MAT/1"])
def test_material_codes_match_existing_catalog_rules(code):
    result = parse_workbook(workbook(code=code))
    assert result.material_rows[0].status == "error"


def test_material_name_matches_catalog_length_and_total_row_bound():
    assert parse_workbook(workbook(name="x" * 161)).material_rows[0].status == "error"
    assert len(parse_workbook(workbook(count=499)).materials) == 499
    with pytest.raises(OperationError):
        parse_workbook(workbook(count=500))


def test_malformed_sheet_xml_is_validation_error():
    content = replace_entry(workbook(), "xl/worksheets/sheet2.xml", lambda _: b"<broken>")
    with pytest.raises(OperationError) as caught:
        parse_workbook(content)
    assert caught.value.status_code == 422


def test_duplicate_zip_entry_and_expansion_limit_rejected():
    for path, payload in [("xl/workbook.xml", b"duplicate"), ("padding.xml", b"x" * 10_485_761)]:
        stream = BytesIO(workbook())
        with ZipFile(stream, "a", ZIP_DEFLATED) as archive:
            if path == "xl/workbook.xml":
                with pytest.warns(UserWarning, match="Duplicate name"):
                    archive.writestr(path, payload)
            else:
                archive.writestr(path, payload)
        with pytest.raises(OperationError):
            parse_workbook(stream.getvalue())


@pytest.mark.parametrize("payload_kind", ["macro", "external", "formula"])
def test_active_excel_content_rejected(payload_kind):
    content = workbook()
    if payload_kind == "macro":
        stream = BytesIO(content)
        with ZipFile(stream, "a", ZIP_DEFLATED) as archive:
            archive.writestr("xl/vbaProject.bin", b"macro")
        content = stream.getvalue()
    elif payload_kind == "external":
        content = replace_entry(
            content,
            "xl/_rels/workbook.xml.rels",
            lambda payload: payload.replace(
                b"<Relationship Id=", b"<Relationship TargetMode = 'External' Id=", 1
            ),
        )
    else:
        content = replace_entry(
            content,
            "xl/worksheets/sheet1.xml",
            lambda payload: (
                payload.decode()
                .replace("</worksheet>", "<f>1+1</f></worksheet>")
                .replace('encoding="UTF-8"', 'encoding="utf-16"')
                .encode("utf-16")
            ),
        )
    with pytest.raises(OperationError, match="unsafe_import_workbook"):
        parse_workbook(content)


def test_template_keeps_catalog_strings_literal_and_identifier_cells_text():
    content = _build_template([("AREA", '=HYPERLINK("https://example.com")')])
    assert not parse_workbook(content).equipment_rows
    book = load_workbook(BytesIO(content))
    try:
        reference = book.worksheets[2]["B2"]
        assert reference.data_type == "s" and reference.hyperlink is None
        assert book.worksheets[0].column_dimensions["A"].number_format == "@"
        assert book.worksheets[1].column_dimensions["A"].number_format == "@"
    finally:
        book.close()
